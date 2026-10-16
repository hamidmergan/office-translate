import zipfile

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from office_translate import translate_file


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def build(path):
    d = docx.Document()
    d.add_heading("Annual report", level=1)
    p = d.add_paragraph("This is ")
    p.add_run("bold").bold = True
    p.add_run(" and this is ")
    p.add_run("italic").italic = True
    p.add_run(".")
    split = d.add_paragraph()
    split.add_run("one ")
    split.add_run("paragraph")
    link = d.add_paragraph("Visit ")
    add_hyperlink(link, "our website", "https://example.org")
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left cell"
    table.cell(0, 1).text = "right cell"
    d.sections[0].header.paragraphs[0].text = "page header"
    d.sections[0].footer.paragraphs[0].text = "page footer"
    d.add_paragraph("12345")
    d.save(path)


def test_docx_text_and_formatting(tmp_path, opts):
    src = tmp_path / "report.docx"
    build(src)
    out = translate_file(src, target_lang="German", **opts)
    assert out == tmp_path / "report.german.docx"

    d = docx.Document(out)
    texts = [p.text for p in d.paragraphs]
    assert "ANNUAL REPORT" in texts
    assert "THIS IS BOLD AND THIS IS ITALIC." in texts
    assert "ONE PARAGRAPH" in texts
    assert "12345" in texts
    runs = d.paragraphs[1].runs
    assert [r.text for r in runs if r.bold] == ["BOLD"]
    assert [r.text for r in runs if r.italic] == ["ITALIC"]
    assert d.tables[0].cell(0, 0).text == "LEFT CELL"
    assert d.tables[0].cell(0, 1).text == "RIGHT CELL"
    assert d.sections[0].header.paragraphs[0].text == "PAGE HEADER"
    assert d.sections[0].footer.paragraphs[0].text == "PAGE FOOTER"
    xml = d.paragraphs[3]._p.xml
    assert "OUR WEBSITE" in xml
    assert "VISIT" in xml


def test_docx_same_format_runs_are_merged(tmp_path, server, opts):
    src = tmp_path / "a.docx"
    build(src)
    translate_file(src, tmp_path / "b.docx", target_lang="German", **opts)
    content = "\n".join(server.user_messages())
    assert "one paragraph" in content


def test_docx_other_parts_unchanged(tmp_path, opts):
    src = tmp_path / "a.docx"
    build(src)
    out = translate_file(src, tmp_path / "b.docx", target_lang="German", **opts)
    with zipfile.ZipFile(src) as a, zipfile.ZipFile(out) as b:
        assert a.namelist() == b.namelist()
        assert a.read("word/styles.xml") == b.read("word/styles.xml")


def test_docx_textbox(tmp_path, opts):
    from helpers import read_part, replace_part

    src = tmp_path / "box.docx"
    d = docx.Document()
    d.add_paragraph("before")
    d.save(src)
    xml = read_part(src, "word/document.xml")
    box = (
        '<w:p><w:r><mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        '<mc:Choice Requires="wps"><w:drawing><wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
        "<w:txbxContent><w:p><w:r><w:t>text in a box</w:t></w:r></w:p></w:txbxContent>"
        "</wp:anchor></w:drawing></mc:Choice></mc:AlternateContent></w:r></w:p>"
    )
    xml = xml.replace("<w:sectPr", box + "<w:sectPr", 1)
    replace_part(src, "word/document.xml", xml.encode("utf-8"))

    out = translate_file(src, tmp_path / "out.docx", target_lang="German", **opts)
    result = read_part(out, "word/document.xml")
    assert "TEXT IN A BOX" in result
    assert "BEFORE" in result


def test_docx_footnotes(tmp_path, opts):
    from helpers import read_part, replace_part

    src = tmp_path / "notes.docx"
    d = docx.Document()
    d.add_paragraph("body")
    d.save(src)
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    notes = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="{w}">'
        '<w:footnote w:id="1"><w:p><w:r><w:t>a footnote</w:t></w:r></w:p></w:footnote></w:footnotes>'
    )
    types = read_part(src, "[Content_Types].xml").replace(
        "</Types>",
        '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-'
        'officedocument.wordprocessingml.footnotes+xml"/></Types>',
    )
    replace_part(src, "[Content_Types].xml", types.encode("utf-8"))
    with zipfile.ZipFile(src, "a") as z:
        z.writestr("word/footnotes.xml", notes)

    out = translate_file(src, tmp_path / "out.docx", target_lang="German", **opts)
    assert "A FOOTNOTE" in read_part(out, "word/footnotes.xml")
