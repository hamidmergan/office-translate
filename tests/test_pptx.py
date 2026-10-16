import pptx
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from helpers import read_part
from office_translate import translate_file


def build(path):
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarterly review"
    body = slide.placeholders[1].text_frame
    body.text = "first point"
    para = body.add_paragraph()
    run = para.add_run()
    run.text = "important"
    run.font.bold = True
    run = para.add_run()
    run.text = " detail"

    box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(1))
    box.text_frame.text = "free text box"

    table = slide.shapes.add_table(1, 2, Inches(1), Inches(6), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "north"
    table.cell(0, 1).text = "south"

    group = slide.shapes.add_group_shape()
    inner = group.shapes.add_textbox(Inches(5), Inches(1), Inches(2), Inches(1))
    inner.text_frame.text = "grouped text"

    slide.notes_slide.notes_text_frame.text = "speaker notes"

    data = CategoryChartData()
    data.categories = ["apples", "pears"]
    data.add_series("sales", (1, 2))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(3), Inches(3), Inches(2), data)
    prs.save(path)


def shape_texts(slide):
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            out.append(shape.text_frame.text)
    return out


def test_pptx(tmp_path, opts):
    src = tmp_path / "deck.pptx"
    build(src)
    out = translate_file(src, target_lang="de", **opts)
    assert out == tmp_path / "deck.de.pptx"

    slide = pptx.Presentation(out).slides[0]
    texts = shape_texts(slide)
    assert "QUARTERLY REVIEW" in texts
    assert "FIRST POINT\nIMPORTANT DETAIL" in texts
    assert "FREE TEXT BOX" in texts

    body = slide.placeholders[1].text_frame.paragraphs[1]
    assert [r.text for r in body.runs if r.font.bold] == ["IMPORTANT"]

    tables = [s.table for s in slide.shapes if s.has_table]
    assert tables[0].cell(0, 0).text == "NORTH"
    groups = [s for s in slide.shapes if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP]
    assert groups[0].shapes[0].text_frame.text == "GROUPED TEXT"
    assert slide.notes_slide.notes_text_frame.text == "SPEAKER NOTES"

    chart = [s for s in slide.shapes if s.has_chart][0].chart
    assert list(chart.plots[0].categories) == ["APPLES", "PEARS"]
    assert chart.series[0].name == "SALES"


def test_layout_placeholders_untouched(tmp_path, opts):
    src = tmp_path / "deck.pptx"
    build(src)
    out = translate_file(src, tmp_path / "out.pptx", target_lang="de", **opts)
    assert read_part(src, "ppt/slideLayouts/slideLayout1.xml") == read_part(out, "ppt/slideLayouts/slideLayout1.xml")
    assert read_part(src, "ppt/slideMasters/slideMaster1.xml") == read_part(out, "ppt/slideMasters/slideMaster1.xml")
